"""Build partners.xlsx from partners_data.py. Run: python3 build_partners_xlsx.py"""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from partners_data import DATE, PARTNERS

COLUMNS = ["Business Name", "Type", "Address", "City", "County", "ZIP", "Phone", "Website", "Notes",
           "Source URL", "Date Collected", "Contacted", "Status"]
COUNTIES = ["Jackson", "Harrison", "George"]
NAVY = PatternFill("solid", fgColor="384A78")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def style(ws, filters=True):
    for cell in ws[1]:
        cell.fill, cell.font = NAVY, WHITE_BOLD
    ws.freeze_panes = "A2"
    if filters:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 3, 60)


def main():
    rows, seen = [], set()
    for name, typ, addr, city, county, zip_, phone, web, src, note in PARTNERS:
        key = re.sub(r"\W", "", name.lower())
        if key in seen:
            continue
        seen.add(key)
        rows.append(dict(zip(COLUMNS, [name, typ, addr, city, county, zip_, phone, web, note, src, DATE])))
    rows.sort(key=lambda r: (r["Type"], COUNTIES.index(r["County"]), r["City"], r["Business Name"].lower()))

    wb = Workbook()
    wb.active.title = "All Partners"
    sheets = [("All Partners", rows)] + [(title, [r for r in rows if r["Type"] == t]) for title, t in
                                         [("Property Managers", "Property Manager"),
                                          ("Insurance Agencies", "Insurance Agency"),
                                          ("Home Builders", "Home Builder")]]
    for title, subset in sheets:
        ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
        ws.append(COLUMNS)
        for r in subset:
            ws.append([r.get(c, "") for c in COLUMNS])
        style(ws)

    ws = wb.create_sheet("Summary")
    ws.append(["Type"] + COUNTIES + ["Total", "With phone"])
    for t in ["Property Manager", "Insurance Agency", "Home Builder"]:
        counts = [sum(1 for r in rows if r["Type"] == t and r["County"] == c) for c in COUNTIES]
        ws.append([t] + counts + [sum(counts), sum(1 for r in rows if r["Type"] == t and r["Phone"])])
    counts = [sum(1 for r in rows if r["County"] == c) for c in COUNTIES]
    ws.append(["Total"] + counts + [len(rows), sum(1 for r in rows if r["Phone"])])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    style(ws, filters=False)

    out = Path(__file__).resolve().parent / "partners.xlsx"
    wb.save(out)
    print(f"{len(rows)} partners, {sum(1 for r in rows if r['Phone'])} with phone -> {out}")


if __name__ == "__main__":
    main()
