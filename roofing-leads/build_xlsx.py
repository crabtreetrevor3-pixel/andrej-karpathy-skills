"""Build leads.xlsx from leads_data.py. Run: python3 build_xlsx.py"""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from leads_data import DATE, LEADS

COLUMNS = ["Business Name", "Category", "Address", "Town", "ZIP", "Phone", "Email", "Email Source URL",
           "Website", "Notes", "Date Collected", "Contacted", "Status"]
TOWNS = ["Pascagoula", "Gautier", "Moss Point"]
NAVY = PatternFill("solid", fgColor="384A78")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def to_rows():
    rows, seen = [], set()
    for name, cat, addr, town, zip_, phone, email, email_src, web, src, note in LEADS:
        key = (re.sub(r"\W", "", name.lower()), re.sub(r"\W", "", addr.lower()))
        if key in seen:
            continue
        seen.add(key)
        notes = [n for n in [note, "" if email else "no email found", f"Source: web search - {src}"] if n]
        rows.append({"Business Name": name, "Category": cat, "Address": addr, "Town": town, "ZIP": zip_,
                     "Phone": phone, "Email": email, "Email Source URL": email_src, "Website": web,
                     "Notes": "; ".join(notes), "Date Collected": DATE})
    return sorted(rows, key=lambda r: (TOWNS.index(r["Town"]), r["Category"], r["Business Name"].lower()))


def style(ws, filters=True):
    for cell in ws[1]:
        cell.fill, cell.font = NAVY, WHITE_BOLD
    ws.freeze_panes = "A2"
    if filters:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 3, 60)


def main():
    rows = to_rows()
    wb = Workbook()
    for title, subset in [("All Leads", rows), ("Phone Only", [r for r in rows if r["Phone"] and not r["Email"]])]:
        ws = wb.active if title == "All Leads" else wb.create_sheet(title)
        ws.title = title
        ws.append(COLUMNS)
        for r in subset:
            ws.append([r.get(c, "") for c in COLUMNS])
        style(ws)

    ws = wb.create_sheet("Summary")
    ws.append(["Category"] + TOWNS + ["Total"])
    for cat in sorted({r["Category"] for r in rows}):
        counts = [sum(1 for r in rows if r["Category"] == cat and r["Town"] == t) for t in TOWNS]
        ws.append([cat] + counts + [sum(counts)])
    totals = [sum(1 for r in rows if r["Town"] == t) for t in TOWNS]
    ws.append(["Total"] + totals + [len(rows)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.append([])
    ws.append(["Contact info"] + TOWNS + ["Total"])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = NAVY, WHITE_BOLD
    for label, test in [("With phone", lambda r: r["Phone"]), ("With email", lambda r: r["Email"]),
                        ("With phone and email", lambda r: r["Phone"] and r["Email"]),
                        ("Phone only (no email)", lambda r: r["Phone"] and not r["Email"]),
                        ("No phone and no email", lambda r: not r["Phone"] and not r["Email"])]:
        counts = [sum(1 for r in rows if r["Town"] == t and test(r)) for t in TOWNS]
        ws.append([label] + counts + [sum(counts)])
    style(ws, filters=False)

    out = Path(__file__).resolve().parent / "leads.xlsx"
    wb.save(out)
    print(f"{len(rows)} businesses, {sum(1 for r in rows if r['Phone'])} with phone, "
          f"{sum(1 for r in rows if r['Email'])} with email -> {out}")


if __name__ == "__main__":
    main()
